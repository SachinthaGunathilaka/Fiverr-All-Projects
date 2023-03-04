from docx.enum.table import WD_TABLE_ALIGNMENT
def createtable(table, a, b):
    t = table
    dataframe = a
    style = b

    if style == str("BASIC"):
        print("BASIC")

        t.allow_autofit = True
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = 'Light Grid Accent 3'

        # add the header rows.
        for j in range(dataframe.shape[-1]):
            t.cell(0, j).text = dataframe.columns[j]

        # add the rest of the data frame
        for i in range(dataframe.shape[0]):
            for j in range(dataframe.shape[-1]):
                t.cell(i + 1, j).text = str(dataframe.values[i, j])

    if style == "NPS":
        print("NPS")

        t.allow_autofit = True
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = 'Light Grid Accent 3'

        # add the header rows.
        for j in range(dataframe.shape[-1]):
            t.cell(0, j).text = dataframe.columns[j]

        # add the rest of the data frame
        for i in range(dataframe.shape[0]):
            for j in range(dataframe.shape[-1]):
                t.cell(i + 1, j).text = str(dataframe.values[i, j])

    if style == "SATS":
        print("SATS")

        t.allow_autofit = True
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = 'Light Grid Accent 3'

        # add the header rows.
        for j in range(dataframe.shape[-1]):
            t.cell(0, j).text = dataframe.columns[j]

        # add the rest of the data frame
        for i in range(dataframe.shape[0]):
            for j in range(dataframe.shape[-1]):
                t.cell(i + 1, j).text = str(dataframe.values[i, j])

    if style == "CROSSED":
        print("CROSSED")

        t.allow_autofit = True
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = 'Light Grid Accent 3'

        # add the header rows.
        for j in range(dataframe.shape[-1]):
            t.cell(0, j).text = dataframe.columns[j]

        # add the rest of the data frame
        for i in range(dataframe.shape[0]):
            for j in range(dataframe.shape[-1]):
                t.cell(i + 1, j).text = str(dataframe.values[i, j])

    return